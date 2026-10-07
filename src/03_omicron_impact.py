"""Task 3 - Do COVID-19 case waves move credit card spending? (2020-21 history -> Omicron read-across)."""
import openpyxl, pandas as pd
import matplotlib.pyplot as plt
from common import DATA, OUT, NAVY, RED, GREY, GREEN, AMBER

wb = openpyxl.load_workbook(DATA / "Task_3-Data.xlsx", data_only=True)
ws = wb["Credit Card Spending"]
hdr = {ws.cell(r, 1).value: r for r in range(1, 12) if ws.cell(r, 1).value}
dates = [ws.cell(hdr["Date"], c).value for c in range(2, ws.max_column + 1)]
cols = [i for i, d in enumerate(dates) if d is not None]
spend = pd.DataFrame({k: [ws.cell(hdr[k], 2 + i).value for i in cols]
                      for k in ["Total Spending", "Delayable", "Social", "Staples", "Work-related"]},
                     index=pd.to_datetime([dates[i] for i in cols])).astype(float)
rows = [r[:2] for r in wb["Covid cases"].iter_rows(min_row=2, values_only=True) if r[0] is not None and r[1] is not None]
cases = pd.Series([r[1] for r in rows], index=pd.to_datetime([r[0] for r in rows])).astype(float)
c7 = cases.rolling(7).mean()

j = spend.join(c7.rename("cases7")).loc["2020-03-01":"2021-12-31"].fillna(0)
corr = j.corr()["cases7"].drop("cases7").round(2); corr.to_csv(OUT / "case_spend_correlation.csv"); print("Correlation with 7-day avg cases, Mar-20 to Dec-21:\n", corr)
full = spend.join(c7.rename("cases7")).loc["2020-01-01":"2021-12-31"].fillna(0).corr()["cases7"]["Total Spending"]
print("Total spending, full 2020-21 incl. Jan-Feb 20:", round(full, 2))

wk = spend.resample("W").mean().loc[:"2021-12-31"]; wc = cases.resample("W").mean().reindex(wk.index).fillna(0)
fig, ax = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
ax[0].plot(wk.index, wk["Total Spending"], color=NAVY); ax[0].set_ylabel("Spending index (Feb-20=100)"); ax[0].axhline(100, color=GREY, lw=.8, ls="--")
a2 = ax[0].twinx(); a2.fill_between(wc.index, wc, color=GREY, alpha=.5); a2.set_ylabel("Avg daily new cases")
for k, col in zip(["Social", "Work-related", "Delayable", "Staples"], [RED, NAVY, AMBER, GREEN]): ax[1].plot(wk.index, wk[k], label=k, color=col)
ax[1].legend(ncol=4, frameon=False); ax[1].set_ylabel("Category index")
ax[0].set_title("Card spending vs COVID-19 cases, 2020-21 (weekly)"); plt.tight_layout(); plt.savefig(OUT / "spending_vs_cases.png", dpi=150)

print("Omicron 7-day peak", c7.idxmax().date(), round(c7.max()), "| 2020-21 peak", round(c7[:"2021-10-31"].max()))
fig, ax = plt.subplots(figsize=(10, 4)); w = cases.resample("W").mean()
ax.bar(w.index, w, width=6, color=[RED if d >= pd.Timestamp("2021-11-01") else GREY for d in w.index]); ax.set_title("Avg daily new cases by week, to 1 May 2022 (Omicron period in red)")
plt.tight_layout(); plt.savefig(OUT / "cases_to_may22.png", dpi=150)
